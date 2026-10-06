"""
Link downloaded markdown files to the bible_contradictions Excel spreadsheet.
Adds two columns: Local File (hyperlink to .md) and Source URL (hyperlink to web).
"""
import os
import re
from pathlib import Path
from copy import copy
import openpyxl
from openpyxl.styles import Font

EXCEL_IN = r"C:\Users\lowes\Downloads\bible_contradictions_HUD_tagged.xlsx"
EXCEL_OUT = r"C:\Users\lowes\Downloads\bible_contradictions_HUD_tagged_linked.xlsx"
MD_DIR = Path("D:/GitHub/crawl4ai/downloaded_pages")


def build_file_index():
    """Build maps from question text and URL to file paths."""
    question_map = {}  # lowercase question -> filepath
    url_map = {}       # source url -> filepath
    all_files = {}     # filename -> filepath

    for f in MD_DIR.glob("*.md"):
        all_files[f.stem.lower()] = f
        text = f.read_text(encoding="utf-8", errors="replace")

        # Extract source URL from YAML header
        for line in text.splitlines()[:5]:
            if line.startswith("url:"):
                url = line[4:].strip()
                url_map[url] = f
                break

        # Extract question from first heading
        for line in text.splitlines():
            cleaned = line.lstrip("#").strip()
            if cleaned and "?" in cleaned and len(cleaned) > 10:
                if not cleaned.startswith("SAB") and not cleaned.startswith("http"):
                    key = cleaned.lower().rstrip("?").strip()
                    question_map[key] = f
                    break

    return question_map, url_map, all_files


def find_file_for_row(question, question_map, all_files):
    """Try to match an Excel question to a downloaded file."""
    if not question:
        return None

    q = question.strip().lower().rstrip("?").strip()

    # Exact match
    if q in question_map:
        return question_map[q]

    # Substring match
    for key, fpath in question_map.items():
        if q in key or key in q:
            return fpath

    # Keyword match on filename
    words = re.findall(r'[a-z]{4,}', q)
    if words:
        best_score = 0
        best_file = None
        for stem, fpath in all_files.items():
            if "contra" not in stem:
                continue
            score = sum(1 for w in words if w in stem)
            if score > best_score and score >= 2:
                best_score = score
                best_file = fpath
        if best_file:
            return best_file

    return None


def find_url_for_file(filepath, url_map):
    """Get the source URL for a file."""
    for url, fpath in url_map.items():
        if fpath == filepath:
            return url
    return None


def main():
    print("Building file index...")
    question_map, url_map, all_files = build_file_index()
    print(f"  {len(question_map)} questions indexed")
    print(f"  {len(url_map)} URLs indexed")
    print(f"  {len(all_files)} total files")

    print(f"\nLoading Excel: {EXCEL_IN}")
    wb = openpyxl.load_workbook(EXCEL_IN)

    # Process each sheet that has a "Question" column
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row < 2:
            continue

        # Find Question column
        q_col = None
        for c in range(1, ws.max_column + 1):
            val = ws.cell(1, c).value
            if val and "question" in str(val).lower():
                q_col = c
                break

        if not q_col:
            continue

        # Add header columns
        link_col = ws.max_column + 1
        url_col = ws.max_column + 2
        ws.cell(1, link_col, "Local File")
        ws.cell(1, link_col).font = Font(bold=True)
        ws.cell(1, url_col, "Source URL")
        ws.cell(1, url_col).font = Font(bold=True)

        matched = 0
        total = 0

        for r in range(2, ws.max_row + 1):
            question = ws.cell(r, q_col).value
            if not question:
                continue
            total += 1

            filepath = find_file_for_row(question, question_map, all_files)
            if filepath:
                matched += 1
                # Local file hyperlink
                file_url = filepath.as_uri()
                cell = ws.cell(r, link_col)
                cell.value = filepath.name
                cell.hyperlink = file_url
                cell.font = Font(color="0000FF", underline="single")

                # Source URL
                source_url = find_url_for_file(filepath, url_map)
                if source_url:
                    ucell = ws.cell(r, url_col)
                    ucell.value = source_url
                    ucell.hyperlink = source_url
                    ucell.font = Font(color="0000FF", underline="single")

        print(f"  {sheet_name}: {matched}/{total} rows linked")

    print(f"\nSaving: {EXCEL_OUT}")
    wb.save(EXCEL_OUT)
    print("Done!")


if __name__ == "__main__":
    main()
