import pandas as pd
import openpyxl
from openpyxl import load_workbook

excel_path = r"C:\Users\lowes\Downloads\UNIFIED_HUD_MASTER_FINAL.xlsx"
print(f"Analyzing: {excel_path}\n")

# Load with openpyxl
wb = load_workbook(excel_path)
print(f"Sheet names: {wb.sheetnames}")

for sheet_idx, sheet_name in enumerate(wb.sheetnames):
    ws = wb[sheet_name]
    print(f"\n=== Sheet {sheet_idx}: '{sheet_name}' ===")
    print(f"Dimensions: {ws.dimensions}")
    
    # Show first 10 rows
    print(f"\nFirst 10 rows (raw values):")
    for row_idx in range(1, min(11, ws.max_row + 1)):
        values = []
        for col_idx in range(1, min(10, ws.max_column + 1)):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is None:
                values.append("(empty)")
            elif len(str(val)) > 50:
                values.append(str(val)[:50] + "...")
            else:
                values.append(str(val))
        print(f"  Row {row_idx}: {values}")

