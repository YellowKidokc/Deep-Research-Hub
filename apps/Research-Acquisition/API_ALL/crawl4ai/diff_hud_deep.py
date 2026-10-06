import os, sys, io
os.environ["PYTHONUTF8"] = "1"
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from openpyxl import load_workbook

NEW_FILE = r"D:\GitHub\crawl4ai\UNIFIED_HUD_TACTICAL_DATABASE_1.xlsx"
OLD_FILE = r"C:\Users\lowes\Downloads\UNIFIED_HUD_TACTICAL_DATABASE.xlsx"

def safe(val, maxlen=80):
    if val is None:
        return "(empty)"
    s = str(val).encode("ascii", errors="replace").decode("ascii")
    if len(s) > maxlen:
        s = s[:maxlen] + "..."
    return s

wb_new = load_workbook(NEW_FILE, read_only=True, data_only=True)
wb_old = load_workbook(OLD_FILE, read_only=True, data_only=True)

print("DEEP CELL-BY-CELL COMPARISON")
print("=" * 80)

total_diffs = 0

for sname in wb_new.sheetnames:
    if sname not in wb_old.sheetnames:
        print("Sheet " + sname + " only in NEW")
        continue
    
    ws_new = wb_new[sname]
    ws_old = wb_old[sname]
    
    new_rows = list(ws_new.iter_rows(values_only=True))
    old_rows = list(ws_old.iter_rows(values_only=True))
    
    max_rows = max(len(new_rows), len(old_rows))
    
    sheet_diffs = 0
    for r in range(max_rows):
        if r >= len(new_rows):
            print("Sheet " + sname + " Row " + str(r+1) + ": exists in OLD but not NEW")
            sheet_diffs += 1
            continue
        if r >= len(old_rows):
            print("Sheet " + sname + " Row " + str(r+1) + ": exists in NEW but not OLD")
            sheet_diffs += 1
            continue
        
        nr = new_rows[r]
        orr = old_rows[r]
        max_cols = max(len(nr), len(orr))
        
        for c in range(max_cols):
            nv = nr[c] if c < len(nr) else None
            ov = orr[c] if c < len(orr) else None
            if str(nv) != str(ov):
                print("DIFF Sheet=" + sname + " Row=" + str(r+1) + " Col=" + str(c+1))
                print("  OLD: " + safe(ov, 120))
                print("  NEW: " + safe(nv, 120))
                sheet_diffs += 1
                total_diffs += 1
    
    if sheet_diffs == 0:
        print("Sheet " + sname + ": IDENTICAL (" + str(len(new_rows)) + " rows)")
    else:
        print("Sheet " + sname + ": " + str(sheet_diffs) + " differences found")

wb_new.close()
wb_old.close()

print("=" * 80)
print("TOTAL DIFFERENCES: " + str(total_diffs))
if total_diffs == 0:
    print("THE TWO FILES ARE IDENTICAL IN CONTENT.")
