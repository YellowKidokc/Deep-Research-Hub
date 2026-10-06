import os, sys, io
os.environ["PYTHONUTF8"] = "1"
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from openpyxl import load_workbook
from collections import Counter

NEW_FILE = r"D:\GitHub\crawl4ai\UNIFIED_HUD_TACTICAL_DATABASE_1.xlsx"
OLD_FILE = r"C:\Users\lowes\Downloads\UNIFIED_HUD_TACTICAL_DATABASE.xlsx"

def safe(val, maxlen=120):
    if val is None:
        return "(empty)"
    s = str(val)
    s = s.encode("ascii", errors="replace").decode("ascii")
    if len(s) > maxlen:
        s = s[:maxlen] + "..."
    return s

def analyze_workbook(filepath, label):
    print("=" * 90)
    print("  ANALYSIS: " + label)
    print("  File: " + filepath)
    print("=" * 90)
    wb = load_workbook(filepath, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    print("[1] SHEET NAMES (" + str(len(sheet_names)) + " total): " + str(sheet_names))
    all_headers = {}
    for sname in sheet_names:
        ws = wb[sname]
        print("-" * 80)
        print("  SHEET: " + sname)
        print("-" * 80)
        rows_data = []
        for row in ws.iter_rows(values_only=True):
            rows_data.append(row)
        total_rows = len(rows_data)
        if total_rows == 0:
            print("  (empty sheet)")
            continue
        headers = list(rows_data[0])
        while headers and headers[-1] is None:
            headers.pop()
        num_cols = len(headers)
        all_headers[sname] = headers
        data_rows = rows_data[1:]
        non_empty_count = sum(1 for r in data_rows if any(cell is not None and str(cell).strip() != "" for cell in r))
        print("  Total rows (incl header): " + str(total_rows))
        print("  Non-empty data rows: " + str(non_empty_count))
        print("  Columns: " + str(num_cols))
        print("  [2] COLUMN HEADERS:")
        for i, h in enumerate(headers, 1):
            print("    Col " + str(i) + ": " + safe(h))
        print("  [3] SAMPLE DATA (rows 2-4):")
        for row_idx in range(min(3, len(data_rows))):
            row = data_rows[row_idx]
            print("    --- Row " + str(row_idx+2) + " ---")
            for col_idx, h in enumerate(headers):
                val = row[col_idx] if col_idx < len(row) else None
                print("      " + safe(h,40) + ": " + safe(val))
        cbc_col_idx = None
        for i, h in enumerate(headers):
            if h and "CBC Rebuttal" in str(h):
                cbc_col_idx = i
                break
        print("  [4] CBC Rebuttal Text COLUMN:")
        if cbc_col_idx is not None:
            print("    FOUND at col " + str(cbc_col_idx+1) + " header=" + safe(headers[cbc_col_idx]))
            populated = sum(1 for r in data_rows if cbc_col_idx < len(r) and r[cbc_col_idx] is not None and str(r[cbc_col_idx]).strip() != "")
            empty_cnt = non_empty_count - populated
            print("    Populated: " + str(populated))
            print("    Empty: " + str(empty_cnt))
            samples = []
            for r in data_rows:
                if cbc_col_idx < len(r) and r[cbc_col_idx] is not None and str(r[cbc_col_idx]).strip() != "":
                    samples.append(safe(r[cbc_col_idx], 200))
                    if len(samples) >= 3:
                        break
            for sv in samples:
                print("    Sample: " + sv)
        else:
            print("    NOT FOUND")
        hud_col_idx = None
        for i, h in enumerate(headers):
            if h and "HUD" in str(h).upper() and "CODE" in str(h).upper():
                hud_col_idx = i
                break
        print("  [5] HUD CODES:")
        if hud_col_idx is not None:
            print("    FOUND at col " + str(hud_col_idx+1) + " header=" + safe(headers[hud_col_idx]))
            hud_vals = [str(r[hud_col_idx]).strip() for r in data_rows if hud_col_idx < len(r) and r[hud_col_idx] is not None and str(r[hud_col_idx]).strip() != ""]
            print("    Assigned: " + str(len(hud_vals)) + "/" + str(non_empty_count))
            dist = Counter(hud_vals)
            print("    Distribution (" + str(len(dist)) + " unique):")
            for code, count in dist.most_common(30):
                print("      " + safe(code,60) + ": " + str(count))
        else:
            print("    NOT FOUND")
        print("  [7] TOTAL NON-EMPTY DATA ROWS: " + str(non_empty_count))
    wb.close()
    return all_headers

def compare_headers(old_h, new_h):
    print("=" * 90)
    print("  COMPARISON: NEW vs OLD")
    print("=" * 90)
    old_s = set(old_h.keys())
    new_s = set(new_h.keys())
    print("  Sheets OLD only: " + str(old_s - new_s if old_s - new_s else "none"))
    print("  Sheets NEW only: " + str(new_s - old_s if new_s - old_s else "none"))
    print("  Sheets in BOTH: " + str(old_s & new_s))
    for sname in sorted(old_s & new_s):
        oc = set(str(h) for h in old_h[sname] if h is not None)
        nc = set(str(h) for h in new_h[sname] if h is not None)
        added = nc - oc
        removed = oc - nc
        print("  Sheet: " + sname)
        print("    Old cols: " + str(len(oc)) + " New cols: " + str(len(nc)))
        if added:
            print("    [6] NEW COLUMNS:")
            for c in sorted(added):
                print("      + " + safe(c))
        else:
            print("    [6] No new columns")
        if removed:
            print("    REMOVED COLUMNS:")
            for c in sorted(removed):
                print("      - " + safe(c))
        else:
            print("    No removed columns")
    for sname in sorted(new_s - old_s):
        print("  NEW SHEET: " + sname)
        print("    Columns: " + str([safe(h) for h in new_h.get(sname,[]) if h]))
    for sname in sorted(old_s - new_s):
        print("  REMOVED SHEET: " + sname)
        print("    Columns: " + str([safe(h) for h in old_h.get(sname,[]) if h]))

print("#" * 90)
print("UNIFIED HUD TACTICAL DATABASE - VERSION COMPARISON")
print("#" * 90)
new_h = analyze_workbook(NEW_FILE, "NEW (_1.xlsx)")
old_h = analyze_workbook(OLD_FILE, "OLD (.xlsx)")
compare_headers(old_h, new_h)
print("#" * 90)
print("END OF REPORT")
print("#" * 90)
