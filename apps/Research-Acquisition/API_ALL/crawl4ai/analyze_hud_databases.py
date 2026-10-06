import os
os.environ["PYTHONUTF8"] = "1"

from openpyxl import load_workbook
from collections import Counter

NEW_FILE = r"D:\GitHub\crawl4ai\UNIFIED_HUD_TACTICAL_DATABASE_1.xlsx"
OLD_FILE = r"C:\Users\lowes\Downloads\UNIFIED_HUD_TACTICAL_DATABASE.xlsx"

def analyze_workbook(filepath, label):
    print("=" * 90)
    print(f"  ANALYSIS: {label}")
    print(f"  File: {filepath}")
    print("=" * 90)

    wb = load_workbook(filepath, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    print(f"\n[1] SHEET NAMES ({len(sheet_names)} total): {sheet_names}")

    all_headers = {}

    for sname in sheet_names:
        ws = wb[sname]
        print(f"\n" + "-" * 80)
        print(f"  SHEET: '{sname}'")
        print("-" * 80)

        rows_data = []
        for row in ws.iter_rows(values_only=True):
            rows_data.append(row)

        total_rows = len(rows_data)
        if total_rows == 0:
            print("  (empty sheet)")
            continue

        headers = list(rows_data[0]) if total_rows > 0 else []
        while headers and headers[-1] is None:
            headers.pop()
        num_cols = len(headers)
        all_headers[sname] = headers

        data_rows = rows_data[1:]
        non_empty_count = 0
        for r in data_rows:
            if any(cell is not None and str(cell).strip() != "" for cell in r):
                non_empty_count += 1

        print(f"  Total rows (incl. header): {total_rows}")
        print(f"  Non-empty data rows:       {non_empty_count}")
        print(f"  Columns:                   {num_cols}")

        print(f"\n  [2] COLUMN HEADERS (row 1):")
        for i, h in enumerate(headers, 1):
            print(f"      Col {i:>3}: {h}")

        print(f"\n  [3] SAMPLE DATA (rows 2-4):")
        for row_idx in range(0, min(3, len(data_rows))):
            row_num = row_idx + 2
            row = data_rows[row_idx]
            print(f"\n      --- Row {row_num} ---")
            for col_idx, h in enumerate(headers):
                val = row[col_idx] if col_idx < len(row) else None
                val_str = str(val) if val is not None else "(empty)"
                if len(val_str) > 120:
                    val_str = val_str[:120] + "..."
                print(f"        {h}: {val_str}")

        cbc_col_idx = None
        for i, h in enumerate(headers):
            if h and "CBC Rebuttal" in str(h):
                cbc_col_idx = i
                break

        print(f"\n  [4] 'CBC Rebuttal Text' COLUMN:")
        if cbc_col_idx is not None:
            print(f"      FOUND at column index {cbc_col_idx + 1} (header: '{headers[cbc_col_idx]}')")
            populated = 0
            empty = 0
            sample_values = []
            for r in data_rows:
                val = r[cbc_col_idx] if cbc_col_idx < len(r) else None
                if val is not None and str(val).strip() != "":
                    populated += 1
                    if len(sample_values) < 3:
                        sv = str(val)
                        if len(sv) > 200:
                            sv = sv[:200] + "..."
                        sample_values.append(sv)
                else:
                    empty += 1
            print(f"      Populated: {populated}")
            print(f"      Empty:     {empty}")
            if sample_values:
                print(f"      Sample values:")
                for sv in sample_values:
                    print(f"        - {sv}")
        else:
            print("      NOT FOUND in this sheet.")

        hud_col_idx = None
        for i, h in enumerate(headers):
            if h and "HUD" in str(h).upper() and "CODE" in str(h).upper():
                hud_col_idx = i
                break
        if hud_col_idx is None:
            for i, h in enumerate(headers):
                if h and "HUD" in str(h).upper():
                    hud_col_idx = i
                    break

        print(f"\n  [5] HUD CODES COLUMN:")
        if hud_col_idx is not None:
            print(f"      FOUND at column index {hud_col_idx + 1} (header: '{headers[hud_col_idx]}')")
            hud_values = []
            for r in data_rows:
                val = r[hud_col_idx] if hud_col_idx < len(r) else None
                if val is not None and str(val).strip() != "":
                    hud_values.append(str(val).strip())
            print(f"      Total assigned: {len(hud_values)} / {non_empty_count} data rows")
            if hud_values:
                dist = Counter(hud_values)
                print(f"      Distribution ({len(dist)} unique values):")
                for code, count in dist.most_common(30):
                    disp = code if len(code) <= 80 else code[:80] + "..."
                    print(f"        {disp}: {count}")
                if len(dist) > 30:
                    print(f"        ... and {len(dist) - 30} more unique values")
        else:
            print("      NOT FOUND in this sheet.")

        print(f"\n  [7] TOTAL NON-EMPTY DATA ROWS: {non_empty_count}")

    wb.close()
    return all_headers


def compare_headers(old_headers, new_headers):
    print("\n" + "=" * 90)
    print("  COMPARISON: NEW vs OLD")
    print("=" * 90)

    old_sheets = set(old_headers.keys())
    new_sheets = set(new_headers.keys())

    sheets_added = new_sheets - old_sheets
    sheets_removed = old_sheets - new_sheets
    sheets_common = old_sheets & new_sheets

    print(f"\n  Sheets in OLD only: {sheets_removed if sheets_removed else '(none)'}")
    print(f"  Sheets in NEW only: {sheets_added if sheets_added else '(none)'}")
    print(f"  Sheets in BOTH:     {sheets_common if sheets_common else '(none)'}")

    for sname in sorted(sheets_common):
        old_cols = set(str(h) for h in old_headers[sname] if h is not None)
        new_cols = set(str(h) for h in new_headers[sname] if h is not None)

        added = new_cols - old_cols
        removed = old_cols - new_cols

        print(f"\n  Sheet '{sname}':")
        print(f"    Old columns: {len(old_cols)}")
        print(f"    New columns: {len(new_cols)}")
        if added:
            print(f"    [6] NEW COLUMNS ADDED:")
            for c in sorted(added):
                print(f"         + {c}")
        else:
            print(f"    [6] No new columns added.")
        if removed:
            print(f"    COLUMNS REMOVED:")
            for c in sorted(removed):
                print(f"         - {c}")
        else:
            print(f"    No columns removed.")


if __name__ == "__main__":
    print("\n" + "#" * 90)
    print("#  UNIFIED HUD TACTICAL DATABASE - VERSION COMPARISON REPORT")
    print("#" * 90)

    new_headers = analyze_workbook(NEW_FILE, "NEW VERSION (UNIFIED_HUD_TACTICAL_DATABASE_1.xlsx)")
    old_headers = analyze_workbook(OLD_FILE, "OLD VERSION (UNIFIED_HUD_TACTICAL_DATABASE.xlsx)")
    compare_headers(old_headers, new_headers)

    print("\n" + "#" * 90)
    print("#  END OF REPORT")
    print("#" * 90)
