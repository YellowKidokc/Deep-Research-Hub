"""
Integrate CARM.org and Defending Inerrancy files into the 6-bucket HUD structure.

These files are SUPPLEMENTARY REFERENCE MATERIAL - they go into a sources/ subfolder
within each bucket. The primary files (matched to Excel rows) stay at the top level.

Also extracts the 10 essays from the docx into 7_ESSAYS/.
"""
import os
import re
import sys
import shutil
from pathlib import Path
from collections import Counter

os.environ["PYTHONUTF8"] = "1"
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

HUD_DIR = Path("D:/GitHub/crawl4ai/bible_contradictions_HUD")
MD_DIR = Path("D:/GitHub/crawl4ai/downloaded_pages")

# ===== AUTO-TAGGER RULES (from hud_tools/auto_tagger.py) =====
RULES = {
    'META': {
        'priority': ['false dilemma', 'argument from silence', 'logical fallacy', 'non-sequitur',
                     'straw man', 'begging the question', 'both true', 'not mutually exclusive',
                     'different senses', 'compatible', 'complementary', 'does not follow',
                     'category error', 'no actual contradiction', 'both can be true',
                     'not a real contradiction', 'misunderstands the text'],
        'secondary': ['fallacious', 'faulty reasoning', 'misreading', 'misunderstanding',
                      'harmonize', 'reconcile', 'both accounts', 'no conflict'],
        'negative': ['manuscript', 'copy error', 'scribe', 'hebrew letter', 'numeral variant'],
    },
    'SCRIBE': {
        'priority': ['copyist error', 'scribal variant', 'numeral corruption', 'manuscript tradition',
                     'textual criticism', 'dalet', 'resh', 'yod', 'vav', 'hebrew letter',
                     'transmission error', 'textual variant', 'copying mistake', 'scribal error'],
        'secondary': ['copy error', 'number discrepancy', 'ancient manuscripts', 'name variant',
                      'minor numerical', 'text critical', 'original reading'],
        'negative': ['context', 'old testament law', 'covenant change', 'audience'],
    },
    'STRIPPING': {
        'priority': ['out of context', 'read the full passage', 'surrounding verses',
                     'incomplete quotation', 'cherry-picking', 'partial quote', 'context shows',
                     'broader passage', 'read in context', 'taken out of context'],
        'secondary': ['missing context', 'when read together', 'in light of', 'full picture',
                      'rest of the chapter', 'preceding verses', 'following verses'],
        'negative': ['manuscript', 'numeral', 'copyist', 'scribal'],
    },
    'VANTAGE': {
        'priority': ['different perspective', 'different witness', 'vantage point',
                     'complementary accounts', 'independent testimony', 'multiple angles',
                     'eyewitness variation', 'different author', 'different gospel writer'],
        'secondary': ['eyewitness', 'angle', 'perspective', 'viewpoint', 'remembered differently',
                      'additional detail', 'saw from', 'reported from'],
        'negative': ['translation', 'hebrew word', 'covenant change', 'copyist'],
    },
    'COVENANT': {
        'priority': ['old covenant', 'new covenant', 'old testament law', 'ceremonial law',
                     'civil law', 'moral law', 'fulfilled in christ', 'hebrews', 'galatians',
                     'law of moses', 'dispensation', 'mosaic law'],
        'secondary': ['old vs new', 'law and grace', 'no longer required', 'fulfilled',
                      'superseded', 'abrogated', 'temporary regulation'],
        'negative': ['manuscript', 'copyist', 'witness perspective', 'translation artifact'],
    },
    'SEMANTIC': {
        'priority': ['hebrew word', 'greek word', 'translation', 'original language',
                     'word range', 'multiple meanings', 'idiom', 'semantic range',
                     'in the original', 'hebrew term', 'greek term'],
        'secondary': ['ben means', 'agape', 'phileo', 'english equivalent', 'linguistic',
                      'literal translation', 'figure of speech', 'hebrew idiom'],
        'negative': ['scribal error', 'manuscript', 'covenant', 'audience adaptation'],
    },
    'DISTINCT': {
        'priority': ['separate event', 'two events', 'happened twice', 'different occasion',
                     'distinct incident', 'not the same event', 'two different', 'occurred twice'],
        'secondary': ['earlier and later', 'first and second time', 'repeated', 'another instance',
                      'similar but distinct', 'two occasions'],
        'negative': ['translation', 'copyist', 'logical fallacy', 'context removal'],
    },
    'ROUNDING': {
        'priority': ['round number', 'approximation', 'approximate figure', 'rounded figure',
                     'rounded number', 'ancient rounding', 'general figure', 'numerical precision'],
        'secondary': ['exact vs round', 'estimate', 'numerical convention', 'rough number',
                      'roughly', 'precision'],
        'negative': ['context', 'translation', 'witness', 'audience'],
    },
    'AUDIENCE': {
        'priority': ['jewish audience', 'roman audience', 'greek audience', 'target audience',
                     'narrative purpose', 'emphasis for readers', 'adapted for', 'writing to'],
        'secondary': ['matthew for jews', 'mark for romans', 'luke for greeks', 'reader oriented',
                      'different emphasis', 'pastoral purpose'],
        'negative': ['copyist', 'manuscript', 'translation artifact', 'scribal'],
    },
    'PHENOM': {
        'priority': ['observational language', 'phenomenological', 'poetic language', 'metaphor',
                     'anthropomorphic', 'literary device', 'figurative', 'poetic expression'],
        'secondary': ['appears to', 'from human perspective', 'manner of speaking', 'poetic',
                      'not literal science', 'figures of speech'],
        'negative': ['copyist', 'manuscript', 'audience adaptation', 'covenant transition'],
    },
}

# 10 HUD codes -> 6 buckets
BUCKET_MAP = {
    "SCRIBE":    "1_SCRIBAL_ERROR",
    "ROUNDING":  "1_SCRIBAL_ERROR",
    "VANTAGE":   "2_POLICE_REPORT",
    "AUDIENCE":  "2_POLICE_REPORT",
    "COVENANT":  "3_TIMELINE_SHIFT",
    "SEMANTIC":  "4_DEFINITION_GAME",
    "DISTINCT":  "4_DEFINITION_GAME",
    "STRIPPING": "5_REPORTER_RULE",
    "META":      "6_MISSING_CONTEXT",
    "PHENOM":    "6_MISSING_CONTEXT",
}


def tag_text(text):
    """Score text against all HUD categories. Returns (best_code, confidence, best_score)."""
    text_lower = text.lower()
    scores = {}

    for code, rules in RULES.items():
        score = 0
        for kw in rules['priority']:
            if kw in text_lower:
                score += 2
        for kw in rules['secondary']:
            if kw in text_lower:
                score += 1
        for kw in rules['negative']:
            if kw in text_lower:
                score -= 3
        scores[code] = score

    ranked = sorted(scores.items(), key=lambda x: x[1], reverse=True)
    best_code = ranked[0][0]
    best_score = ranked[0][1]
    second_score = ranked[1][1] if len(ranked) > 1 else 0

    if best_score >= 4 and best_score > second_score + 2:
        confidence = 'HIGH'
    elif best_score >= 2:
        confidence = 'MEDIUM'
    elif best_score >= 1:
        confidence = 'LOW'
    else:
        confidence = 'NONE'
        best_code = 'META'  # Default to META (6_MISSING_CONTEXT) for unclassifiable

    return best_code, confidence, best_score


def integrate_source_files():
    """Tag and copy CARM + Defending Inerrancy files into bucket sources/ folders."""

    # Clean and recreate sources/ subfolder in each bucket
    for bucket_name in set(BUCKET_MAP.values()):
        sources_dir = HUD_DIR / bucket_name / "sources"
        if sources_dir.exists():
            shutil.rmtree(sources_dir)
        sources_dir.mkdir(exist_ok=True)

    # Identify files to process (CARM + Defending Inerrancy only)
    carm_files = sorted(MD_DIR.glob("carm_org_*.md"))
    di_files = sorted(MD_DIR.glob("defendinginerrancy_*.md"))

    all_files = []
    for f in carm_files:
        all_files.append(("CARM", f))
    for f in di_files:
        all_files.append(("DI", f))

    total = len(all_files)
    print(f"Files to integrate: {len(carm_files)} CARM + {len(di_files)} DI = {total}")

    stats = Counter()
    conf_stats = Counter()
    bucket_counts = Counter()

    for i, (source, filepath) in enumerate(all_files, 1):
        try:
            text = filepath.read_text(encoding="utf-8", errors="replace")
        except Exception as e:
            print(f"  [{i:>4}/{total}] ERR  {filepath.name} ({e})")
            stats["error"] += 1
            continue

        # Use first 5000 chars for tagging (enough for classification, fast)
        sample = text[:5000]
        code, confidence, score = tag_text(sample)
        bucket = BUCKET_MAP.get(code, "6_MISSING_CONTEXT")

        # Copy to bucket sources/
        dest = HUD_DIR / bucket / "sources" / filepath.name
        shutil.copy2(filepath, dest)

        stats["ok"] += 1
        conf_stats[confidence] += 1
        bucket_counts[bucket] += 1

        if i % 100 == 0 or i == total:
            print(f"  [{i:>4}/{total}] Processed... ({stats['ok']} ok, {stats.get('error', 0)} err)")

    return stats, conf_stats, bucket_counts


def extract_essays():
    """Extract the 10 essays from the docx into 7_ESSAYS/."""
    docx_path = Path("D:/GitHub/crawl4ai/categorical_demolition_10_essays.docx")
    essays_dir = HUD_DIR / "7_ESSAYS"

    if not docx_path.exists():
        print(f"\nDocx not found: {docx_path}")
        return 0

    try:
        from docx import Document
    except ImportError:
        print("\npython-docx not installed. Installing...")
        import subprocess
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
        from docx import Document

    doc = Document(str(docx_path))

    # Extract all text with paragraph styles
    full_text = []
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else "Normal"
        full_text.append((style_name, para.text))

    # Split into essays by detecting "Essay N:" or "ESSAY N" patterns
    essays = []
    current_essay = None
    current_lines = []

    for style, text in full_text:
        # Detect essay boundary
        essay_match = re.match(r'(?:Essay|ESSAY)\s*(\d+)[:\s]*(.*)', text.strip(), re.IGNORECASE)
        heading_match = (style and 'Heading' in str(style) and essay_match)

        if essay_match and (heading_match or len(text.strip()) < 200):
            # Save previous essay
            if current_essay is not None and current_lines:
                essays.append((current_essay, current_lines))
            current_essay = int(essay_match.group(1))
            title = essay_match.group(2).strip(' -–—:')
            current_lines = [f"# Essay {current_essay}: {title}\n"]
        elif current_essay is not None:
            # Add to current essay
            if style and 'Heading' in style:
                level = 2
                if '2' in style:
                    level = 2
                elif '3' in style:
                    level = 3
                current_lines.append(f"\n{'#' * level} {text}\n")
            else:
                current_lines.append(text)

    # Save last essay
    if current_essay is not None and current_lines:
        essays.append((current_essay, current_lines))

    if not essays:
        # Fallback: split by any "Essay N" pattern in the text
        all_text = "\n".join(t for _, t in full_text)
        parts = re.split(r'(?=(?:Essay|ESSAY)\s*\d+[:\s])', all_text)
        for part in parts:
            m = re.match(r'(?:Essay|ESSAY)\s*(\d+)[:\s]*(.*?)[\n]', part, re.IGNORECASE)
            if m:
                num = int(m.group(1))
                title = m.group(2).strip(' -–—:')
                content = f"# Essay {num}: {title}\n\n{part[m.end():]}"
                essays.append((num, [content]))

    # Essay number -> bucket mapping
    ESSAY_BUCKET = {
        1: "6_MISSING_CONTEXT",   # META
        2: "1_SCRIBAL_ERROR",     # SCRIBE
        3: "5_REPORTER_RULE",     # STRIPPING
        4: "2_POLICE_REPORT",     # VANTAGE
        5: "3_TIMELINE_SHIFT",    # COVENANT
        6: "4_DEFINITION_GAME",   # SEMANTIC
        7: "4_DEFINITION_GAME",   # DISTINCT
        8: "1_SCRIBAL_ERROR",     # ROUNDING
        9: "2_POLICE_REPORT",     # AUDIENCE
        10: "6_MISSING_CONTEXT",  # PHENOM
    }

    print(f"\nExtracted {len(essays)} essays from docx")

    for num, lines in essays:
        content = "\n".join(lines) if isinstance(lines[0], str) and not lines[0].startswith("#") else "\n".join(lines)

        # Save to 7_ESSAYS/
        essay_file = essays_dir / f"Essay_{num:02d}.md"
        essay_file.write_text(content, encoding="utf-8")

        # Also copy to the relevant bucket
        bucket = ESSAY_BUCKET.get(num)
        if bucket:
            bucket_essay = HUD_DIR / bucket / f"_ESSAY_{num:02d}.md"
            bucket_essay.write_text(content, encoding="utf-8")

        print(f"  Essay {num}: {len(content):,} chars -> {essay_file.name}" +
              (f" + {bucket}" if bucket else ""))

    return len(essays)


def print_final_summary():
    """Print the final state of all folders."""
    print(f"\n{'='*60}")
    print("FINAL HUD STRUCTURE")
    print(f"{'='*60}")

    total_primary = 0
    total_sources = 0
    total_essays = 0

    for d in sorted(HUD_DIR.iterdir()):
        if not d.is_dir():
            continue

        primary = len([f for f in d.glob("*.md") if f.name != "README.md"])
        sources_dir = d / "sources"
        sources = len(list(sources_dir.glob("*.md"))) if sources_dir.exists() else 0

        if d.name == "7_ESSAYS":
            total_essays = primary
            print(f"  {d.name}/  ({primary} essays)")
        elif d.name == "8_NEEDS_INDIVIDUAL_WORK":
            print(f"  {d.name}/  ({primary} placeholders)")
        else:
            total_primary += primary
            total_sources += sources
            print(f"  {d.name}/  ({primary} primary + {sources} source refs)")

    print(f"\n  TOTALS:")
    print(f"    Primary contradiction files: {total_primary}")
    print(f"    Supplementary source files:  {total_sources}")
    print(f"    Essays:                      {total_essays}")
    print(f"    Grand total:                 {total_primary + total_sources + total_essays}")


def main():
    print("=" * 60)
    print("INTEGRATING CARM + DEFENDING INERRANCY INTO HUD BUCKETS")
    print("=" * 60)

    # Step 1: Tag and integrate source files
    stats, conf_stats, bucket_counts = integrate_source_files()

    print(f"\n--- Tagging Results ---")
    print(f"  Processed: {stats['ok']}, Errors: {stats.get('error', 0)}")
    print(f"  Confidence: {dict(conf_stats)}")
    print(f"  Per bucket: {dict(bucket_counts)}")

    # Step 2: Extract essays
    essay_count = extract_essays()

    # Step 3: Final summary
    print_final_summary()

    print("\nDone!")


if __name__ == "__main__":
    main()
