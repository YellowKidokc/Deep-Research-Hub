from __future__ import annotations
import html,json,shutil
from pathlib import Path
from .paths import API_HOME

def latest_runs(paper:Path):
 for station in sorted((paper/'02_RUNS').glob('*')):
  runs=sorted((p for p in station.iterdir() if p.is_dir()),reverse=True) if station.is_dir() else []
  if runs: yield station.name,runs[0]
def build(paper:Path)->tuple[Path,Path]:
 report=paper/'03_REPORT'; report.mkdir(parents=True,exist_ok=True); stats_path=report/'statistics.json'; stats=json.loads(stats_path.read_text(encoding='utf-8')) if stats_path.exists() else {'metrics':[],'headline':[]}
 prototype=(API_HOME/'templates/statistics_matrix.html').read_text(encoding='utf-8')
 prototype=prototype.replace('DEMO NUMBERS · generated to show the layout, not a real run','LIVE RUN DATA · unavailable metrics remain explicitly provisional').replace('Every number on this page is generated demo data.','Values supplied by statistics.json; fields not yet emitted by the configured metric suite remain provisional.')
 payload=json.dumps(stats,ensure_ascii=False).replace('</','<\\/')
 bridge='''\n/* live statistics.json bridge */\nvar LIVE_DATA=__PAYLOAD__;\nvar liveByName={};(LIVE_DATA.metrics||[]).forEach(function(x){liveByName[(x.id||x.name||'').toLowerCase().replace(/_/g,' ')]=x;});\nMETRICS.forEach(function(m){var x=liveByName[m.name.toLowerCase()];if(x&&typeof x.value==='number'){m.val=x.value;m.src=x.method==='AI-judged'?'A':'L';m.ga=x.academic_percentile==null?null:x.academic_percentile;m.gc=x.corpus_percentile==null?50:x.corpus_percentile;m.series=x.series_percentile==null?50:x.series_percentile;m.dv=x.change_since_previous||0;m.low=!!x.runs_disagree;}});\n'''.replace('__PAYLOAD__',payload)
 prototype=prototype.replace('/* ---------- wiring ---------- */',bridge+'\n/* ---------- wiring ---------- */')
 sections=[]
 for label,run in latest_runs(paper):
  page=run/f'{label}.html'
  if page.exists(): sections.append(f'<section data-station="{html.escape(label)}">{page.read_text(encoding="utf-8",errors="replace")}</section>')
 combined=prototype.replace('</body>','<section class="panel"><h2>Station outputs</h2>'+''.join(sections)+'</section></body>') if '</body>' in prototype else prototype+''.join(sections)
 html_path=report/'report.html'; html_path.write_text(combined,encoding='utf-8')
 xlsx=report/'report.xlsx'
 import importlib.util
 if importlib.util.find_spec('openpyxl'):
  from openpyxl import Workbook
  wb=Workbook(); wb.remove(wb.active)
  for label,run in latest_runs(paper):
   ws=wb.create_sheet(label[:31]); source=run/f'{label}.json'
   obj=json.loads(source.read_text(encoding='utf-8')) if source.exists() else {}
   ws.append(['JSON']); ws.append([json.dumps(obj,ensure_ascii=False)])
  if not wb.worksheets: wb.create_sheet('README').append(['No station outputs found'])
  wb.save(xlsx)
 else: xlsx.write_text('Install openpyxl to create a native workbook.\n',encoding='utf-8')
 return html_path,xlsx
